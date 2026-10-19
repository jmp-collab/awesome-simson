import json, html
from openpyxl import load_workbook
from top10 import TOP
E=html.escape
S=json.load(open('data/scored.json'))
wb=load_workbook('out/02_Neunkirchen_Leads_Scoring.xlsx',data_only=True)['Top 10 Angebote']
vals=[r[13] for r in wb.iter_rows(min_row=2,max_row=13,values_only=True)]
pipe,real=vals[10],vals[11]
fmt=lambda n:f"{n:,.0f}".replace(",",".")
tagc={"Umsatzhebel":"hebel","Quick Win":"quick","Solide Chance":"solid"}
jump="".join(f'<a href="#{t["k"]}">{i}. {E(t["name"].split(",")[0])}</a>' for i,t in enumerate(TOP,1))
leads=[]
for i,t in enumerate(TOP,1):
    pre=f"img/{i:02d}_{t['k']}"
    leads.append(f'''<article class="lead" id="{t["k"]}">
<div class="lhead"><div><div class="nr">{i:02d}</div><h2>{E(t["name"])}</h2>
<div class="meta">{E(t["branche"])} · {E(t["addr"])}, 66538 Neunkirchen · <span class="tel">{E(t["tel"])}</span></div></div>
<div class="tags"><span class="tag {tagc.get(t["label"],"")}">{E(t["label"])}</span><span class="tag">Google {E(t["g"])}</span><span class="tag">Besuch: {E(t["visit"])}</span></div></div>
<div class="cmp-tools"><span class="seg" role="group" aria-label="Ansicht"><button type="button" data-view="desk" aria-pressed="true">Desktop</button><button type="button" data-view="mob" aria-pressed="false">Handy</button></span><span class="lbl">Schieber ziehen: links Vorher, rechts Nachher</span></div>
<div class="cmp desk"><img class="before" src="{pre}_vorher_desk.jpg" data-desk="{pre}_vorher_desk.jpg" data-mob="{pre}_vorher_mob.jpg" alt="Aktuelle Website {E(t["name"])}"><img class="after" src="{pre}_nachher_desk.jpg" data-desk="{pre}_nachher_desk.jpg" data-mob="{pre}_nachher_mob.jpg" alt="Entwurf neue Website {E(t["name"])}"><span class="cap l">VORHER</span><span class="cap r">NACHHER</span><div class="bar"></div><div class="knob">⇆</div><input id="r-{t["k"]}" type="range" min="0" max="100" value="50" aria-label="Vorher Nachher Schieber"></div>
<div class="cols"><div class="box"><h3>Was dem Inhaber auffällt</h3><ul>{"".join(f"<li>{E(a)}</li>" for a in t["audit"])}</ul><p class="hebel">{E(t["hebel"])}</p></div>
<div class="box"><h3>Angebot</h3><div class="price"><div><span>Basis-Relaunch</span><b>{E(t["p"][0])}</b></div><div class="rec"><span>Empfohlen, einmalig</span><b>{E(t["p"][1])}</b></div><div><span>Betreuung / Monat</span><b>{E(t["p"][2])}</b></div></div>
<p style="font-size:15px;color:var(--mut)">{E(t["paket"])}</p><h3>Branchen-Extras</h3><ul>{"".join(f"<li>{E(s)}</li>" for s in t["svc"])}</ul></div></div>
<div class="talk"><div class="lbl">Gesprächsleitfaden</div><div class="say"><q>{E(t["open"])}</q><button type="button">Kopieren</button></div><p><b>Einwand:</b> {E(t["obj"])}</p></div>
</article>''')
top={ "Schlüssel Thielen","Asia Wok Bistro","Pizza Heimservice Gondoliere","Beauty Nails Ornella","Krankengymnastik Rutz","Markt Apotheke","Wirtshaus Römer Steakhaus","Damaskus Haus","Tanzwerkstatt Susanne Sprung","Hotel Zum Ellenfeld"}
rank="".join(f'<tr class="{"top" if r["name"] in top else ""}"><td class="n">{r["rank"]}</td><td>{E(r["name"])}</td><td>{E(r["branch"])}</td><td class="n">{(str(r["g_avg"]).replace(".",",")+" ("+str(r["g_n"])+")") if r["g_avg"] else "-"}</td><td class="n">{r["A"]}</td><td class="n">{r["B"]}</td><td class="n">{r["C"]}</td><td class="n">{r["D"]}</td><td class="n">{r["total"]}<span class="meter"><i style="width:{r["total"]}%"></i></span></td><td>{E(r["label"])}</td></tr>' for r in S[:30])
M=[("Döner, Kebab, Imbiss",["Eigenes Bestellsystem ohne Provision, Abholung und Lieferung: 990 bis 1.990 €","Mehrsprachige Speisekarte (DE, EN, TR oder AR): 190 bis 390 €","QR-Aufsteller 'Bewerte uns' und 'Nächstes Mal direkt bestellen': 99 bis 199 €","TikTok und Instagram Reels vom Spieß, 3 pro Woche: 390 bis 790 €/Monat","Bewertungsantworten mit KI-Entwurf zur Freigabe: 79 bis 149 €/Monat"],"Hebel: Plattformprovision 23 bis 30 % laut Anbieterangabe. Orient Grill (Rang 6, 616 Bewertungen) leitet auf unverschlüsseltes HTTP um: nächster Kandidat."),
("Pizzeria und Lieferdienst",["Bestellsystem plus Liefergebiet und Mindestbestellwert: 1.490 bis 2.490 €","Digitale Stempelkarte per QR auf dem Karton: 149 bis 349 €","Karriereseite plus Google-Beitrag für Fahrer und Küchenhilfe: 190 bis 390 €","Lieferzeit-Hinweis und Bewertungsmanagement: 79 bis 149 €/Monat"],"La Campanella, Da Enzo und Piccobello laufen ebenfalls auf Plattform-Standardseiten."),
("Restaurant, Steakhaus, Wirtshaus",["Online-Reservierung: 290 bis 590 €","Feiern- und Gruppenseite für die Weihnachtssaison: 290 bis 490 €","Online-Gutscheinverkauf: 290 bis 590 €","Englische Kurzkarte: 149 bis 290 €","Social Media 2 bis 3 Beiträge pro Woche: 290 bis 590 €/Monat"],"Oktober ist der richtige Monat: Weihnachtsfeiern werden jetzt gebucht."),
("Hotel und Pension",["Zimmeranfrage oder Direktbuchung: 390 bis 990 €","Pflegbarer Wochenplan (Stammessen) mit Google-Beitrag: 190 bis 390 €","Raumvermietung als Landingpage: 190 bis 390 €","Zimmer-Fotoshooting: 290 bis 490 €"],"Pension am Stummplatz (Jimdo, Google 3,5 bei 157 Bewertungen) braucht zuerst Bewertungsmanagement."),
("Spielothek und Spielhalle",["Sachliche Infoseite: Öffnungszeiten, Anfahrt, Ausstattung, Jugend- und Spielerschutz: 590 bis 990 €","Google-Profil-Pflege mit Ambientefotos: 49 bis 99 €/Monat","Karriereseite für Servicekräfte: 190 bis 390 €","Bewertungsantworten: 49 bis 99 €/Monat"],"Recht: Glücksspielstaatsvertrag 2021 und saarländisches Spielhallenrecht. Keine Gewinnversprechen, keine Ansprache Minderjähriger, Texte juristisch prüfen lassen. Lux Play (Rang 56) hat eine neue Baukastenseite, Ansatz dort eher Google-Profil und Rechtssicherheit der Texte."),
("Apotheke",["E-Rezept-Anleitung, Vorbestellung, Botendienst prominent: in Relaunch enthalten","Notdienst- und Öffnungsstatus live: 149 bis 290 €","Aktionen als Google-Beitrag: 49 bis 99 €/Monat","Karriereseite PTA: 190 bis 390 €"],"Heilmittelwerbegesetz beachten. Bären und Linden Apotheke antworten nur unverschlüsselt: weitere Kandidaten."),
("Physiotherapie, Zahnarzt, Praxen",["Terminanfrage und Hausbesuch-Formular: 190 bis 390 €","Karriereseite mit Stellenportal-Verteilung: 290 bis 590 €","Leistungsseiten verständlich erklärt: 90 bis 150 € je Seite","Barrierearme Gestaltung: im Relaunch enthalten"],"Heilmittelwerbegesetz und Berufsordnung: keine Werbung mit Patientendank, keine Heilversprechen."),
("Tanzschule, Fitness, Kampfsport",["Filterbarer Kursplan plus Probestunden-Buchung: 390 bis 890 €","Aufführungs- oder Event-Landingpage mit Tickets: 290 bis 590 €","Reels aus dem Training, 2 bis 3 pro Woche: 290 bis 590 €/Monat","Ferienkurs-Kampagne: 190 bis 390 €"],"Bei Kindern immer schriftliche Einwilligung der Eltern für Fotos und Videos."),
("Beauty, Nagelstudio, Friseur",["Online-Terminbuchung mit Erinnerung: 290 bis 590 €","Google-Profil neu anlegen und verifizieren: 190 bis 290 €","Instagram mit Vorher/Nachher der Arbeiten: 290 bis 590 €/Monat","Gutscheinverkauf online: 290 bis 590 €"],"Seiler Friseurteam und Kosmetik Rova antworten nur unverschlüsselt: weitere Kandidaten."),
("Handwerk, Schlüsseldienst, Service",["Klick-zum-Anrufen und Notruf prominent: im Relaunch enthalten","Google-Profil als Hauptkanal: 190 bis 390 € einmalig","Einträge in Verzeichnissen vereinheitlichen (NAP): 149 bis 290 €","Lokale Backlinks: Vereine, Stadtportal, Lokalpresse: 290 bis 690 €"],"Für Notdienste zählt die Google-Suche fast allein."),
("Kanzleien und Steuerberatung",["Rechtsgebiete oder Leistungen als eigene Seiten: 90 bis 150 € je Seite","Mandanten-Upload oder Portal-Link: 190 bis 390 €","Karriereseite: 290 bis 590 €"],"Langläufer: mehrere Partner entscheiden. Berufsrecht beachten. Arminia (Rang 15) ist optisch der älteste Auftritt in dieser Gruppe.")]
matrix="".join(f'<div class="c"><h3>{E(h)}</h3><ul>{"".join(f"<li>{E(x)}</li>" for x in L)}</ul><p class="warn">{E(w)}</p></div>' for h,L,w in M)
P=[("Onepager-Relaunch","Mobil, HTTPS, Impressum, Datenschutz, Kontakt, Google-Grundsetup, KI-gestützte Texte aus vorhandenem Inhalt","590 bis 1.190 €"),
("Mehrseitige Website","4 bis 8 Unterseiten, Team, Karriere, Leistungen","1.190 bis 2.490 €"),
("Gastro mit Bestell- oder Reservierungssystem","Digitale Speisekarte, eigener Bestellweg, Google-Menü","1.490 bis 3.490 €"),
("Google-Unternehmensprofil","Einmalig optimieren / laufend pflegen","190 bis 390 € / 49 bis 149 € pro Monat"),
("Hosting und Wartung","Updates, Backups, bis 30 Minuten Änderungen pro Monat","29 bis 69 € pro Monat"),
("Social Media","Redaktionsplan, 2 bis 3 Beiträge pro Woche, KI-Untertitel","290 bis 890 € pro Monat"),
("Fotoshooting vor Ort","Laden, Produkte, Team, 20 bearbeitete Bilder","190 bis 490 €")]
prices="".join(f"<tr><td><b>{E(a)}</b></td><td>{E(b)}</td><td class=\"mono\">{E(c)}</td></tr>" for a,b,c in P)
D=[("Tag 1 · Mittwoch",[("10:00","Schlüssel Thielen, Hüttenbergstraße 30"),("10:45","Physiotherapie Rutz, Max-Braun-Straße 1"),("11:15","Markt Apotheke, gleiches Haus, nach Frau Daniels fragen"),("14:45","Asia Wok Bistro, Brückenstraße 26"),("15:30","Damaskus Haus, Bahnhofstraße 39"),("16:15","Wirtshaus Römer, Bahnhofstraße 5 (dienstags Ruhetag)")]),
("Tag 2 · Donnerstag",[("09:30","Beauty Nails Ornella anrufen, Termin vereinbaren"),("14:00","Hotel Zum Ellenfeld, Zweibrücker Straße 14"),("15:00","Tanzwerkstatt, Hüttenbergstraße 6"),("16:00","Gondoliere, Storchenplatz 1 (öffnet 17 Uhr, montags zu)"),("17:00","Gegencheck nur-HTTP-Seiten am Handy: Orient Grill, Seiler, Rova, Bella Italia")])]
plan="".join(f'<div class="day"><h3>{E(h)}</h3>{"".join(f"<div class=\"slot\"><b>{a}</b><span>{E(b)}</span></div>" for a,b in L)}</div>' for h,L in D)
s=open('page_template.html').read()
for k,v in {"__PIPE__":fmt(pipe),"__REAL__":fmt(real),"__JUMP__":jump,"__LEADS__":"\n".join(leads),"__RANK__":rank,"__MATRIX__":matrix,"__PRICES__":prices,"__PLAN__":plan}.items(): s=s.replace(k,v)
open('site/index.html','w').write(s)
print(len(s), pipe, real)
