import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from outreach_text import SEG, TIP
from vorlagen import VORLAGEN, BETREFF

A = json.load(open("data/all_v2.json"))
IMP = {}
for f in ("data/imp_old.json", "data/imp_b12.json", "data/imp_b3.json"):
    for x in json.load(open(f)): IMP[x["name"]] = x
TOP10 = {"Schlüssel Thielen","Asia Wok Bistro","Pizza Heimservice Gondoliere","Beauty Nails Ornella","Krankengymnastik Rutz","Markt Apotheke","Wirtshaus Römer Steakhaus","Damaskus Haus","Tanzwerkstatt Susanne Sprung","Hotel Zum Ellenfeld"}
SEGNAME = dict(liefer="Lieferdienst/Imbiss", gastro="Gastronomie", lebensmittel="Lebensmittelhandwerk", hotel="Hotel/Pension", apotheke="Apotheke", praxis="Praxis/Gesundheit", beauty="Beauty/Körperpflege", freizeit="Freizeit/Sport/Kurse", spielhalle="Spielhalle", kfz="Kfz/Zweirad", handwerk="Handwerk/Bau", beratung="Beratung/Kanzlei", bestatter="Bestatter", handel="Handel/Dienstleistung")

DISPLAY = {"Didié Hermann Zeichen- und Bürobedarf":"Hermann Didié Bürobedarf","Heil Schlosserei GmbH":"Schlosserei Heil","Ratzky GmbH Bedachungen":"Ratzky Bedachungen",
"Schneider Birgit Musikunterricht":"Musikunterricht Birgit Schneider","Sascha Kieczewsky Planungsbüro":"Planungsbüro Kieczewsky","Regitz Kreativdesign":"Regitz Creativ Design",
"Orthopädie Lattrich GmbH Wolfgang":"Orthopädie Lattrich","Elektro Hobler UG (haftungsbeschränkt) Elektroinstallationen":"Elektro Hobler","VSV Hochbau GmbH":"VSV Hochbau",
"DEJA Vu Second Hand Exclusive":"Déjà Vu Second Hand","Koble-Tassone Jeannine Friseur":"Hair-Design Koble","Radiologie am Boxberg, Dr. med. M. Nebel, Dr. med. J. Klinkner":"Radiologie am Boxberg",
"Klaproth Holger Dr.med.":"Praxis Dr. Klaproth","Schmidt Thomas Michael Dr.med. u. Schmitt Kerstin Dr.med.":"Hausarztpraxis Wellesweiler","König Bernd Dipl.-Ing. Architekt":"Architekt Bernd König",
"Hegi S. Fenstervertrieb u. Montage":"Fenster Hegi","Decker Fenster Haustüren Rollladen GmbH":"Decker Fenster","LOGOPÄDIE Krämer Monika Schwochow Nathalie":"Logopädie Krämer & Schwochow",
"Jochum Günther Dr.med. Ganzheitliche Medizin":"Praxis Dr. Jochum","Müller Pierrette Dr. Psychotherapeutische Praxis":"Praxis Dr. Pierrette Müller","Meyer Walter , Meyer Christel Yogadozenten":"Yogazentrum Meyer",
"Kirschhock Bernd Haus Montage Service":"Haus Montage Service Kirschhock","Lembcke Melanie Praxis für Osteopathie":"Osteopathie Lembcke",
"Schneider Oliver Dr.med. u. Schemmann Frank Dr.med. Chirurgische Gemeinschaftspraxis":"Chirurgische Praxis am Boxberg","Tennis-Gastronomie TuS 1860 (Brigitte Faber)":"Tennis-Gastronomie TuS 1860",
"HTB Haus-Technik Bechtel GmbH":"HTB Haus-Technik Bechtel","Bierstubb (Dieter Rinner)":"Bierstubb","Jochum Günther Dr. med. Internisten-Kardiologie":"Privatpraxis Dr. Groß",
"Naturheilpraxis Olga Bernhardt Heilpraktikerin":"Naturheilpraxis Bernhardt","Angresius Dr. Arztpraxis für Frauenheilkunde und Geburtshilfe":"Frauenarztpraxis Dr. Angresius",
"Müller Philippe Dres.med. Gemeinschaftspraxis für Neurologie und Psychiatrie , Chrysostomos Kyprianou":"Praxis Müller & Kyprianou","Lorig":"Geflügel Schlemmerlädchen Lorig",
"Spies Susanne Heilpraktikerin":"Naturheilpraxis Susanne Spies","Derouet Harry Prof. Dr.med. Pönicke Christoph Dr.med. Becker Frank Dr.med.":"Urologie Neunkirchen","Bücher König Buchhandel":"Bücher König",
"Hoffmann Reinhard Waffen- u. Sporthaus":"Waffen- und Sporthaus Hoffmann","Intersport Klees Sportfachgeschäft":"Intersport Klees","Grunders Hofgut Menschenhaus":"Hofgut Menschenhaus",
"Störr Christian Rollladenbetrieb":"Störr Fenster","Leibrock Karl Bauer":"Leibrocks Hofladen","Hofmann Viktor Fliesenverlegung":"Fliesen Hofmann","Krankengymnastik Rutz":"Physiotherapie Rutz",
"Pizza Heimservice Gondoliere":"Pizzeria Gondoliere","Wirtshaus Römer Steakhaus":"Wirtshaus Römer","MDM MetallDesignMüller / Innovativer Metallbau / Inh.: Christian Müller-Metallbaumeister":"MetallDesign Müller",
"Schreiner Barbara Fotostudio":"Fotostudio Barbara Schreiner","Dolomiti Pizzeria, Gelateria & Pasticceria":"Dolomiti","Nieland Julia Pasta":"Pastaria by Julie","Asia Haus Thaimassagen":"Asia Haus Thaimassage",
"Lackierservice Forster":"Lackierservice Forster","DIVNA Medizinische Fußpflege":"Podologische Praxis Divna","Seiler Friseurteam":"Friseurteam Seiler"}
def disp(n):
    if n in DISPLAY: return DISPLAY[n]
    x = re.sub(r"\s*\(haftungsbeschränkt\)|\s+(GmbH & Co\.? KG|GmbH|UG|e\.\s?K\.?|e\.Kfm\.?|KG|GbR|PartGmbB|oHG|OHG)\b\.?", "", n)
    x = re.sub(r"\s*(Inh\.|Inhaber).*$", "", x).strip(" ,/-")
    return x
WEAK = ("Auf Ihrer Website gibt es keinen", "Auf Ihrer Startseite habe ich keine")
PS_DEFAULT = "Die Vorschau bleibt auch dann kostenlos, wenn Sie am Ende bei Ihrer jetzigen Seite bleiben."
def anrede(r):
    o = r["owner"]; raw = IMP.get(r["name"], {}).get("owner", "")
    if not o: return "Guten Tag,"
    dr = "Dr. " if re.search(r"\bDr\b\.?", raw) else ""
    last = o.split()[-1]
    if r["gender"] == "f": return f"Guten Tag Frau {dr}{last},"
    if r["gender"] == "m": return f"Guten Tag Herr {dr}{last},"
    return f"Guten Tag {dr}{o},"

def domain(r):
    return re.sub(r"^https?://(www\.)?", "", r["url"]).split("/")[0].lower()

wb = Workbook(); AR = "Arial"
H = Font(name=AR, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3A5F"); N = Font(name=AR, size=10)
INP = PatternFill("solid", fgColor="FFF4CC")
def head(ws, cols, widths, freeze="C2"):
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(1, i); cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = freeze; ws.row_dimensions[1].height = 34

# ---------- Start / Einstellungen
st = wb.active; st.title = "Start"
st.column_dimensions["A"].width = 30; st.column_dimensions["B"].width = 70; st.column_dimensions["C"].width = 60
st["A1"] = "Akquise Neunkirchen, Version 2 (Stand 19.10.2026)"; st["A1"].font = Font(name=AR, bold=True, size=14)
st["A2"] = "Gelbe Felder ausfüllen. Alle Brief- und Mailtexte im Blatt 'Outreach' setzen sich daraus automatisch zusammen."; st["A2"].font = Font(name=AR, italic=True, size=10)
SETTINGS = [
 ("Absender", "Jean", "Ihr voller Name, so wie er unter Brief und Mail stehen soll"),
 ("Telefon", "", "Mobilnummer für Anruf und WhatsApp"),
 ("Absender_Web", "", "Eigene Website oder Vorschau-Domain, nicht die Arbeitgeber-Domain"),
 ("Absender_Mail", "", "Eigene Geschäftsadresse, nicht die Arbeitgeber-Adresse"),
 ("Vorschau_Link", "", "Standard-Link zur Vorschau, falls im Outreach-Blatt kein eigener Link steht"),
 ("Termin_Vorschlag", "Donnerstag gegen 15 Uhr", "Konkreter Terminvorschlag für M1, M2, W1"),
 ("Referenzpreis_Datum", "30.11.2026", "Bis wann der Referenzpreis gilt (F3)"),
 ("Bewertungslink", "", "Kurzlink zu Ihrem Google-Profil zum Bewerten (R1)"),
 ("Empfehlungsbonus", "einem Gutschein über 50 €", "Dankeschön für Empfehlungen (R1)"),
]
for i, (k, v, h) in enumerate(SETTINGS, 4):
    st.cell(i, 1, k).font = Font(name=AR, bold=True, size=10)
    c = st.cell(i, 2, v); c.fill = INP; c.font = N
    st.cell(i, 3, h).font = Font(name=AR, size=9, italic=True, color="555555")
    wb.defined_names[f"S_{k}"] = DefinedName(f"S_{k}", attr_text=f"Start!$B${i}")
r0 = 4 + len(SETTINGS) + 1
for j, line in enumerate([
 "So nutzen Sie die Datei:",
 "1. Blatt 'Outreach' nach Spalte 'Welle' filtern. Welle 1 sind die stärksten Leads mit belegbarem Aufhänger.",
 "2. Erstkontakt per Brief (Spalte 'Brief-Text') oder persönlich. Vor dem Versand die Website einmal am Handy öffnen und den Aufhänger kurz gegenprüfen.",
 "3. E-Mail und WhatsApp erst, wenn der Betrieb zugestimmt hat: im Gespräch, per Antwort auf den Brief oder per Rückruf. Ohne Einwilligung ist Werbung per E-Mail auch an Firmen nach § 7 Abs. 2 UWG unzulässig und abmahnfähig.",
 "4. Status, Datum und Kanal pflegen. Spalte 'Nächster Schritt fällig' rechnet die Nachfass-Termine (Tag 3, 7, 14) selbst aus.",
 "5. Texte ändern Sie zentral im Blatt 'Vorlagen'. Platzhalter in geschweiften Klammern bleiben stehen.",
]):
    st.cell(r0 + j, 1, line).font = Font(name=AR, size=10, bold=(j == 0))

# ---------- Vorlagen
vs = wb.create_sheet("Vorlagen")
head(vs, ["ID", "Name", "Kanal und Rechtsrahmen", "Zeitpunkt", "Betreff", "Text"], [6, 34, 34, 18, 44, 110], freeze="B2")
for v in VORLAGEN:
    vs.append(list(v))
for row in vs.iter_rows(min_row=2):
    for c in row: c.font = N; c.alignment = Alignment(wrap_text=True, vertical="top")
    row[4].fill = INP; row[5].fill = INP
VROW = {v[0]: i for i, v in enumerate(VORLAGEN, 2)}

# ---------- Outreach (Merge)
ow = wb.create_sheet("Outreach")
cols = ["Rang", "Betrieb", "Welle", "Segment", "Inhaber/in", "Anrede", "Straße, PLZ Ort", "Telefon", "E-Mail (nur nach Einwilligung)", "Domain",
        "Aufhänger 1", "Aufhänger 2 (P.S.)", "Winkel", "Extra", "Preis", "Monat", "Ziel", "Tipp", "Vorlage Brief", "Vorschau-Link (optional)",
        "Betreff Brief", "Brief-Text", "Mail nach Besuch (M1)", "Nachfassen Tag 3 (F1)", "Nachfassen Tag 7 (F2)", "Nachfassen Tag 14 (F3)", "WhatsApp (W1)",
        "Kanal", "Datum Erstkontakt", "Status", "Nächster Schritt fällig", "Notiz", "Score", "Label"]
head(ow, cols, [6, 30, 7, 18, 22, 26, 34, 15, 30, 26, 60, 60, 50, 40, 10, 10, 18, 50, 8, 22, 34, 70, 70, 50, 50, 50, 50, 12, 13, 16, 14, 30, 7, 14], freeze="C2")
act = sorted([r for r in A if not r["excluded"]], key=lambda r: r["rank"])
PH = ["Anrede", "Betrieb", "Domain", "Hook1", "Hook2", "Winkel", "Extra", "Preis", "Monat", "Ziel", "Tipp"]
COLOF = {"Anrede": "F", "Betrieb": "B", "Domain": "J", "Hook1": "K", "Hook2": "L", "Winkel": "M", "Extra": "N", "Preis": "O", "Monat": "P", "Ziel": "Q", "Tipp": "R"}
def merge(src, i):
    f = src
    for p in PH:
        f = f'SUBSTITUTE({f},"{{{p}}}",${COLOF[p]}{i})'
    f = f'SUBSTITUTE({f},"{{Vorschau_Link}}",IF($T{i}="",S_Vorschau_Link,$T{i}))'
    for k in ("Absender", "Absender_Web", "Telefon", "Termin_Vorschlag", "Referenzpreis_Datum"):
        f = f'SUBSTITUTE({f},"{{{k}}}",S_{k})'
    return "=" + f
def welle(r):
    strong = r["hook1"] and not r["hook1"].startswith(("Auf Ihrer Website gibt es keinen", "Auf Ihrer Startseite habe ich keine"))
    if r["total"] >= 60 and strong: return 1
    if r["total"] >= 50 and strong: return 2
    return 3
for r in act:
    seg = r["seg"]; w = SEG[seg]
    h1 = r["hook1"] or "Ihre Website zeigt noch nicht, was Ihren Betrieb ausmacht."
    h2 = r["hook2"] if r["hook2"] and r["hook2"] != h1 and not r["hook2"].startswith(WEAK) else PS_DEFAULT
    addr = r["addr"].replace(", 66", "\n66") if r["addr"] else ""
    ow.append([r["rank"], disp(r["name"]), welle(r), SEGNAME[seg], r["owner"], anrede(r), r["addr"], r["tel"], r["email"], domain(r),
               h1, h2, w[0], w[1], w[2], w[3], w[4], TIP[seg], "B2" if r["name"] in TOP10 else "B1", "",
               None, None, None, None, None, None, None, "", None, "offen", None, "", r["total"], r["label"]])
    i = ow.max_row
    ow[f"U{i}"] = merge(f'INDEX(Vorlagen!$E:$E,MATCH($S{i},Vorlagen!$A:$A,0))', i)
    ow[f"V{i}"] = merge(f'INDEX(Vorlagen!$F:$F,MATCH($S{i},Vorlagen!$A:$A,0))', i)
    for col, vid in (("W", "M1"), ("X", "F1"), ("Y", "F2"), ("Z", "F3"), ("AA", "W1")):
        ow[f"{col}{i}"] = merge(f"Vorlagen!$F${VROW[vid]}", i)
    ow[f"AE{i}"] = f'=IF(OR($AC{i}="",$AD{i}="Auftrag",$AD{i}="Abgesagt"),"",IF($AD{i}="Brief versendet",$AC{i}+5,IF($AD{i}="Vorschau gesendet",$AC{i}+3,IF($AD{i}="Nachgefasst 1",$AC{i}+7,IF($AD{i}="Nachgefasst 2",$AC{i}+14,"")))))'
    ow[f"AC{i}"].number_format = "DD.MM.YYYY"; ow[f"AE{i}"].number_format = "DD.MM.YYYY"
for row in ow.iter_rows(min_row=2):
    for c in row: c.font = N; c.alignment = Alignment(wrap_text=True, vertical="top")
for col in ("T", "AB", "AC", "AD", "AF"):
    for c in ow[col][1:]: c.fill = INP
last = ow.max_row
dv = DataValidation(type="list", formula1='"offen,Brief versendet,Besucht,Vorschau gesendet,Nachgefasst 1,Nachgefasst 2,Termin,Angebot,Auftrag,Später,Abgesagt"', allow_blank=True)
ow.add_data_validation(dv); dv.add(f"AD2:AD{last}")
dv2 = DataValidation(type="list", formula1='"Brief,Besuch,E-Mail,WhatsApp,Telefon (Rückruf)"', allow_blank=True)
ow.add_data_validation(dv2); dv2.add(f"AB2:AB{last}")
dv3 = DataValidation(type="list", formula1='"B1,B2"', allow_blank=False)
ow.add_data_validation(dv3); dv3.add(f"S2:S{last}")
ow.conditional_formatting.add(f"AE2:AE{last}", FormulaRule(formula=[f'AND($AE2<>"",$AE2<=TODAY())'], fill=PatternFill("solid", fgColor="F4B6B6")))
ow.auto_filter.ref = f"A1:AH{last}"
for i in range(2, last + 1): ow.row_dimensions[i].height = 60

# ---------- Gesamtliste
gl = wb.create_sheet("Gesamtliste")
gc = ["Rang", "Betrieb", "Branche", "Segment", "Adresse", "Telefon", "Website", "Prüfstatus", "HTTPS", "Mobil", "Ladezeit s", "CMS", "Letzte Änderung", "Facebook", "Instagram", "Strukt. Daten", "Buchung/Bestellung", "Google Ø", "Google Anzahl", "A Modernisierung (0-40)", "B Einfachheit (0-20)", "C Potenzial (0-25)", "D Lücken (0-15)", "Gesamt", "Label", "Liste", "Inhaber/in", "E-Mail", "Impressum"]
head(gl, gc, [6, 32, 26, 18, 34, 15, 34, 12, 7, 7, 8, 16, 22, 8, 8, 8, 9, 8, 8, 10, 10, 10, 10, 8, 14, 8, 22, 28, 34])
SK = dict(live="live", http="nur HTTP", bot="Bot-Schutz", defekt="Fehlerseite")
for r in act:
    mob = "ja" if r["live"] and r.get("viewport") and (r.get("mob_scrollW") or 390) <= 400 else ("nein" if r["live"] else "?")
    gl.append([r["rank"], r["name"], r["branch"], SEGNAME[r["seg"]], r["addr"], r["tel"], r.get("final") or r["url"], SK[r["status_kind"]],
               "ja" if r.get("https") and r["live"] else "nein", mob, r.get("load_s"), (r.get("generator") or "")[:30], r.get("lastmod", ""),
               "ja" if r["fb"] else "nein", "ja" if r["ig"] else "nein", "ja" if r.get("localbiz") else "nein", "ja" if r.get("booking") else "nein",
               r.get("g_avg"), r.get("g_n"), r["A"], r["B"], r["C"], r["D"], None, r["label"], r["quelle"], r["owner"], r["email"], r.get("imp_url", "")])
    i = gl.max_row; gl[f"X{i}"] = f"=SUM(T{i}:W{i})"
for row in gl.iter_rows(min_row=2):
    for c in row: c.font = N
t = Table(displayName="Gesamtliste", ref=f"A1:{get_column_letter(len(gc))}{gl.max_row}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); gl.add_table(t)

# ---------- Ohne Website
nw = wb.create_sheet("Ohne Website")
head(nw, ["Betrieb", "Branche", "Adresse", "Online-Präsenz", "Telefon", "Hinweis", "Betreff (B3)", "Brief-Text (B3)", "Status"], [34, 24, 34, 50, 16, 34, 40, 80, 14])
OLD_NW = [("Eiscafé Venezia","Eiscafé","Stummplatz 3, 66538 Neunkirchen","facebook.com/nkstummplatz3","",""),("Café Sick","Café","Oberer Markt 3, 66538 Neunkirchen","Facebook-Seite","",""),("Tanja's Backhäusle","Bäckerei","Oberer Markt 15, 66538 Neunkirchen","Facebook-Seite","",""),("Pilsstubb Gasthaus","Kneipe","Langenstrichstr. 47, 66538 Neunkirchen","Facebook-Seite","06821 6363595","Google 4,5 (50)"),("Treffpunkt Gaststätte","Kneipe","Zweibrücker Str. 72, 66538 Neunkirchen","Facebook-Seite","",""),("The Sushi Twins","Restaurant Sushi","Bliesstr. 44, 66538 Neunkirchen","facebook.com/TheSushiTwins","",""),("Café Löwe","Café","Hebbelstr. 1, 66538 Neunkirchen","cafe-loewe.eatbu.com","",""),("Armando's Brasserie","Restaurant","Stummplatz 2, 66538 Neunkirchen","nur Branchenprofil","","")]
for x in OLD_NW: nw.append(list(x))
for l in open("data/ohne_website_v2.txt", encoding="utf-8"):
    p = [x.strip() for x in l.rstrip("\n").split("|")]
    if len(p) != 6: continue
    hint = "Alte Domain nicht mehr erreichbar" if any(k in p[4].lower() for k in ("dasoertliche", "11880", "speisekarte")) else ""
    nw.append([p[0], p[1], f"{p[2]}, {p[3]} Neunkirchen".strip(", "), p[4], p[5], hint])
def m3(src, i):
    f = src
    f = f'SUBSTITUTE({f},"{{Anrede}}","Guten Tag,")'
    f = f'SUBSTITUTE({f},"{{Betrieb}}",$A{i})'
    for k in ("Absender", "Absender_Web", "Telefon"):
        f = f'SUBSTITUTE({f},"{{{k}}}",S_{k})'
    return "=" + f
for i in range(2, nw.max_row + 1):
    nw[f"G{i}"] = m3(f"Vorlagen!$E${VROW['B3']}", i)
    nw[f"H{i}"] = m3(f"Vorlagen!$F${VROW['B3']}", i)
    nw[f"I{i}"] = "offen"; nw[f"I{i}"].fill = INP
    nw.row_dimensions[i].height = 60
for row in nw.iter_rows(min_row=2):
    for c in row: c.font = N; c.alignment = Alignment(wrap_text=True, vertical="top")

# ---------- Ausgeschlossen
ex = wb.create_sheet("Ausgeschlossen")
head(ex, ["Betrieb", "Branche", "Website", "Grund"], [36, 26, 40, 40])
for r in sorted([r for r in A if r["excluded"]], key=lambda r: r["excluded"]):
    ex.append([r["name"], r["branch"], r["url"], r["excluded"]])
for row in ex.iter_rows(min_row=2):
    for c in row: c.font = N

# ---------- Betreffzeilen
bz = wb.create_sheet("Betreffzeilen")
head(bz, ["Variante", "Betreff", "Einsatz"], [18, 50, 70], freeze="A2")
for b in BETREFF: bz.append(list(b))
for row in bz.iter_rows(min_row=2):
    for c in row: c.font = N

# ---------- KPIs
kp = wb.create_sheet("KPIs")
kp.column_dimensions["A"].width = 34; kp.column_dimensions["B"].width = 12; kp.column_dimensions["C"].width = 14; kp.column_dimensions["D"].width = 60
kp.append(["Kennzahl", "Anzahl", "Quote", "Formel / Annahme"])
for c in kp[1]: c.font = H; c.fill = HF
R = f"Outreach!$AD$2:$AD${last}"
rows = [
 ("Leads gesamt", f"=COUNTA(Outreach!$B$2:$B${last})", None, ""),
 ("Kontaktiert (Status nicht 'offen')", f'=B2-COUNTIF({R},"offen")', "=IF(B2=0,\"\",B3/B2)", "Anteil der Leads, die angesprochen wurden"),
 ("Reaktion (Vorschau gesendet oder weiter)", f'=COUNTIF({R},"Vorschau gesendet")+COUNTIF({R},"Nachgefasst 1")+COUNTIF({R},"Nachgefasst 2")+COUNTIF({R},"Termin")+COUNTIF({R},"Angebot")+COUNTIF({R},"Auftrag")', "=IF(B3=0,\"\",B4/B3)", "Antwortquote auf Brief oder Besuch"),
 ("Termin vereinbart", f'=COUNTIF({R},"Termin")+COUNTIF({R},"Angebot")+COUNTIF({R},"Auftrag")', "=IF(B4=0,\"\",B5/B4)", "Termin pro Reaktion"),
 ("Angebot erstellt", f'=COUNTIF({R},"Angebot")+COUNTIF({R},"Auftrag")', "=IF(B5=0,\"\",B6/B5)", "Angebot pro Termin"),
 ("Auftrag", f'=COUNTIF({R},"Auftrag")', "=IF(B6=0,\"\",B7/B6)", "Abschluss pro Angebot"),
 ("Abschluss pro Kontakt", "=B7", "=IF(B3=0,\"\",B7/B3)", "Kennzahl für die Planung der nächsten Welle"),
]
for x in rows: kp.append(list(x))
for c in kp["C"][1:]: c.number_format = "0.0%"
for row in kp.iter_rows(min_row=2):
    for c in row: c.font = N

# ---------- Methodik
me = wb.create_sheet("Methodik")
me.column_dimensions["A"].width = 150
for line in [
 "Methodik und Quellen (Stand 19.10.2026)", "",
 "Suchraum: Stadt Neunkirchen (Saar) mit PLZ 66538 (Innenstadt), 66539 (Wellesweiler, Furpach, Kohlhof, Ludwigsthal) und 66540 (Wiebelskirchen, Hangard, Münchwies).",
 "Quellen: Das Örtliche (Themen-Seiten, rund 400 Branchen-Slugs), Gelbe Seiten, 11880, Speisekarte.de, Vereinsverzeichnisse. Domains per DNS und Browser-Abruf geprüft.",
 "Technik: Headless Chromium, Desktop 1366x850 und Mobil 390x844, 19.10.2026 (Liste 2) bzw. 09.10.2026 (Liste 1). Gemessen: HTTPS, Viewport, Breite mobil, Ladezeit, CMS, Last-Modified, Social-Links, schema.org, Buchungs- und Bestellhinweise.",
 "Visueller Score: manuelle Einschätzung jedes Screenshots nach festen Ankern 0 bis 25 (siehe Briefing Anhang A).",
 "Score: A Modernisierung 0-40, B Einfachheit 0-20, C Potenzial 0-25 (Branche, Lage, Google-Nachfrage), D Digitale Lücken 0-15.",
 "Google-Profile: Places-Abfrage für 53 Kandidaten. Leere Felder bedeuten: nicht abgefragt oder nicht eindeutig zuordenbar.",
 "Kontakte: E-Mail und Inhaber automatisch aus Impressum bzw. Kontaktseite extrahiert. Anrede nur mit Frau/Herr, wenn das Impressum es eindeutig zeigt.",
 "Aufhänger: nur Beobachtungen, die aus der Messung oder dem Screenshot belegbar sind. Vor Versand einmal am Handy gegenprüfen, Websites ändern sich.",
 "Rechtsrahmen: Werbung per E-Mail, Fax, SMS und Messenger braucht vorherige ausdrückliche Einwilligung, auch gegenüber Firmen (§ 7 Abs. 2 UWG). Werbebriefe per Post sind zulässig, solange kein Widerspruch vorliegt. Telefonwerbung an Firmen nur bei mutmaßlicher Einwilligung, für Webdesign-Angebote in der Regel nicht gegeben. Keine Rechtsberatung.",
]:
    me.append([line])
me["A1"].font = Font(name=AR, bold=True, size=13)
for c in me["A"][1:]: c.font = N; c.alignment = Alignment(wrap_text=True)

wb.move_sheet("Outreach", offset=-(len(wb.sheetnames) - 2))
wb.save("out/04_Neunkirchen_Leads_Outreach_v2.xlsx")
print("ok", last - 1, "Leads", sum(1 for r in act if welle(r) == 1), "Welle1", sum(1 for r in act if welle(r) == 2), "Welle2")
