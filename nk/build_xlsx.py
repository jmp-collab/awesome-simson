import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment
from top10 import TOP
S=json.load(open('data/scored.json'))
mapn={"thielen":"Schlüssel Thielen","asia":"Asia Wok Bistro","gondoliere":"Pizza Heimservice Gondoliere","ornella":"Beauty Nails Ornella","rutz":"Krankengymnastik Rutz","marktapo":"Markt Apotheke","roemer":"Wirtshaus Römer Steakhaus","damaskus":"Damaskus Haus","tanzwerkstatt":"Tanzwerkstatt Susanne Sprung","ellenfeld":"Hotel Zum Ellenfeld"}
topnames={v:i for i,v in enumerate(mapn.values(),1)}
toplab={mapn[t['k']]:t['label'] for t in TOP}
wb=Workbook(); A="Arial"
wb.properties.creator="Jean"; wb.properties.lastModifiedBy="Jean"
H=Font(name=A,bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F3A5F"); N=Font(name=A,size=10)
def head(ws,cols,widths):
    ws.append(cols)
    for i,c in enumerate(cols,1):
        cell=ws.cell(1,i); cell.font=H; cell.fill=HF; cell.alignment=Alignment(wrap_text=True,vertical="center")
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1]
    ws.freeze_panes="C2"; ws.row_dimensions[1].height=32
ws=wb.active; ws.title="Gesamtliste"
cols=["Rang","Name","Branche","Adresse","Telefon","Website (gelistet)","Website (nach Weiterleitung)","Prüfstatus","HTTPS","Mobil-tauglich","Ladezeit s","CMS / Baukasten","Letzte Änderung (Server)","Facebook verlinkt","Instagram verlinkt","Strukturierte Daten","Bestellung/Buchung","Google Ø","Google Anzahl","A Modernisierung (0-40)","B Einfachheit (0-20)","C Potenzial (0-25)","D Digitale Lücken (0-15)","Gesamtscore","Label","Top 10","Hauptargument","Vertriebsstatus","Nächster Schritt"]
head(ws,cols,[6,30,26,34,15,34,34,20,8,10,9,18,22,10,10,10,11,9,9,11,11,11,11,10,15,7,60,14,18])
for r in S:
    ws.append([r['rank'],r['name'],r['branch'],r['addr'],r['tel'],r['url'],r['final'],r['status'],r['https'],r['mobile'],r['load'],r['cms'],r['lastmod'],r['fb'],r['ig'],r['schema'],r['booking'],r['g_avg'],r['g_n'],r['A'],r['B'],r['C'],r['D'],None,toplab.get(r['name'],r['label']),topnames.get(r['name'],""),r['note'],"offen",""])
    i=ws.max_row; ws.cell(i,24).value=f"=SUM(T{i}:W{i})"
for row in ws.iter_rows(min_row=2):
    for c in row: c.font=N
    if row[25].value:
        for c in row: c.fill=PatternFill("solid",fgColor="FFF4CC")
ws.cell(1,20).comment=Comment("Visueller Score 0-25 nach Ankern (Briefing Anhang A) +8 ohne HTTPS +5 nicht mobil +3 Ladezeit >8s +5 Platzhalter/Default-Seite, max 40. Nicht erreichbare Seiten: Basis 14.","Analyse")
ws.cell(1,18).comment=Comment("Google Places Abfrage 09.10.2026. Leer = nicht abgefragt oder nicht eindeutig zuordenbar.","Analyse")
tab=Table(displayName="Gesamtliste",ref=f"A1:{get_column_letter(len(cols))}{ws.max_row}"); tab.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True); ws.add_table(tab)
w2=wb.create_sheet("Top 10 Angebote")
c2=["Nr","Betrieb","Branche","Adresse","Telefon","Google","Label","Basis von €","Basis bis €","Empfohlen von €","Empfohlen bis €","Monatlich von €","Monatlich bis €","Jahreswert Empfohlen (Mitte) €","Empfohlenes Paket","Branchen-Services","Audit-Beleg 1","Audit-Beleg 2","Audit-Beleg 3","Hebel","Gesprächseinstieg","Einwand + Antwort","Beste Besuchszeit"]
head(w2,c2,[5,30,26,24,14,16,13,10,10,11,11,11,11,14,50,60,50,50,50,60,60,60,30])
num=lambda s:[int(x.replace('.','')) for x in re.findall(r'[\d.]+',s)]
for i,t in enumerate(TOP,1):
    b=num(t['p'][0]); e=num(t['p'][1]); m=num(t['p'][2])
    w2.append([i,t['name'],t['branche'],t['addr']+", 66538 Neunkirchen",t['tel'],t['g'],t['label'],b[0],b[1],e[0],e[1],m[0],m[1],None,t['paket'],"\n".join("- "+s for s in t['svc']),*t['audit'],t['hebel'],t['open'],t['obj'],t['visit']])
    r=w2.max_row; w2.cell(r,14).value=f"=(J{r}+K{r})/2+12*(L{r}+M{r})/2"
r=w2.max_row+1
w2.cell(r,2).value="Summe Pipeline (Mitte, wenn alle 10 zusagen)"; w2.cell(r,14).value=f"=SUM(N2:N{r-1})"
w2.cell(r+1,2).value="Realistisch bei 30 % Abschlussquote"; w2.cell(r+1,14).value=f"=N{r}*0.3"
w2.cell(r+2,2).value="Annahme Abschlussquote 30 % = Hypothese für persönliche Ansprache mit fertiger Vorschau; nach 20 Besuchen mit echten Zahlen ersetzen."
for row in w2.iter_rows(min_row=2):
    for c in row: c.font=Font(name=A,size=10); c.alignment=Alignment(wrap_text=True,vertical="top")
for rr in (r,r+1): w2.cell(rr,2).font=Font(name=A,bold=True); w2.cell(rr,14).font=Font(name=A,bold=True)
for col in "HIJKLMN":
    for c in w2[col][1:]: c.number_format='#,##0" €"'
w3=wb.create_sheet("Leads ohne Website")
head(w3,["Name","Branche","Adresse","Online-Präsenz","Hinweis"],[34,26,36,60,50])
for x in [("Eiscafé Venezia","Eiscafé","Stummplatz 3","facebook.com/nkstummplatz3",""),("Café Sick","Café","Oberer Markt 3","Facebook-Seite",""),("Tanja's Backhäusle","Bäckerei","Oberer Markt 15","Facebook-Seite",""),("Pilsstubb Gasthaus","Kneipe","Langenstrichstr. 47","Facebook-Seite",""),("Treffpunkt Gaststätte","Kneipe","Zweibrücker Str. 72","Facebook-Seite",""),("The Sushi Twins","Restaurant Sushi","Bliesstr. 44","facebook.com/TheSushiTwins",""),("Café Löwe","Café","Hebbelstr. 1","cafe-loewe.eatbu.com (Baukasten-Subdomain)",""),("Eiscafé Dolomiti","Eiscafé/Pizzeria","Oberer Markt 10","Metro-Baukasten + Liefersoft-Shop",""),("Armando's Brasserie","Restaurant","Stummplatz 2","nur Branchenprofil sellwerk.de",""),("aqua bar lounge club","Bar/Club","Jahnstr. 2","aqua-nk.de zeigt Plesk-Standardseite","Domain vorhanden, Website fehlt; Google-Profil ohne Bewertungen, Betrieb vor Besuch prüfen")]:
    w3.append([x[0],x[1],x[2]+", 66538 Neunkirchen",x[3],x[4]])
w4=wb.create_sheet("Ausgeschlossen (Ketten)")
head(w4,["Name","Grund"],[40,50])
for n in ["Fair Play Casino","Tipico","Burger King","McDonald's Getrey","KFC","Subway","Pizza Hut","Nordsee","Ditsch","Tchibo","Globus Restaurant","McFit","clever fit","Holiday Inn Express","Barbarossa Bäckerei","Bäckerei Ziegler","Bäckerei Gillen","Fielmann, Apollo, dm, Rossmann, Deichmann, Douglas, Thalia, Klier, Christ, Kraemer, Abele, Amplifon, KIND, Blume 2000, essanelle, DERTOUR"]:
    w4.append([n,"Kette oder Filialist mit zentral gesteuerter Website"])
w5=wb.create_sheet("Methodik")
for line in ["Methodik und Quellen (Stand 09.10.2026)","",
"1. Betriebe: Das Örtliche, Gelbe Seiten, 11880 (PLZ 66538), Domains per HTTP-Abruf geprüft. 95 Betriebe mit eigener Domain, 10 Leads ohne Website, Ketten ausgeschlossen.",
"2. Technik: Headless Chromium, Desktop 1366x850 und Mobil 390x844. Gemessen: HTTPS, Viewport, horizontale Breite mobil, Ladezeit bis Load-Event, Generator/CMS, jQuery, Last-Modified, Social-Links, schema.org, Bestell-/Buchungshinweise.",
"3. Nicht prüfbar: Seiten, die nur unverschlüsselt (HTTP) antworten oder auf HTTP umleiten, waren in der Prüfumgebung nicht abrufbar. Das ist selbst ein Befund (Browser zeigt 'Nicht sicher'), vor dem Besuch aber kurz am Handy gegenprüfen.",
"4. Google: Places-Abfrage für 28 Kandidaten (Schnitt, Anzahl, Öffnungszeiten). Bewertungstexte nur zusammengefasst.",
"5. Score: A Modernisierung 0-40, B Einfachheit 0-20, C Geschäftspotenzial 0-25, D Digitale Lücken 0-15. Gesamtscore = Formel. Visueller Teil von A ist eine manuelle Einschätzung nach festen Ankern.",
"6. Top 10: Score plus Branchenmix, Begehbarkeit (Ladengeschäft Innenstadt) und Belegbarkeit (Screenshot vorhanden).",
"7. Preise: Marktannahmen für Kleinbetriebe im Saarland, keine erhobenen Wettbewerberpreise. Vor dem ersten Angebot mit 2 bis 3 Angeboten lokaler Agenturen plausibilisieren.",
"8. Plattformprovision: Anbieterangabe 23 bis 30 % pro Marktplatzbestellung (get-sides.de, Artikel vom 07.10.2026, Quelle mit Eigeninteresse). Echten Vertrag beim Inhaber erfragen.",
"Legende Top-10-Blatt: Preisspalten H bis M sind Eingaben (anpassbar), Spalte N wird berechnet."]:
    w5.append([line])
w5.column_dimensions['A'].width=150; w5['A1'].font=Font(name=A,bold=True,size=13)
for c in w5['A'][1:]: c.font=Font(name=A,size=10)
wb.save('out/02_Neunkirchen_Leads_Scoring.xlsx')
